# worker/pipeline/document_processor.py
"""
Layer 5 - Document/Text Processor.

Extracts native text from PDF/DOCX/TXT where a real text layer exists.
Flags documents with no extractable text (scanned images) so they route
to OCR instead of wasting effort here. Chunks long text for downstream
model input.
"""

from dataclasses import dataclass
from pathlib import Path

import pdfplumber
from docx import Document as DocxDocument


# Most captioning/summarization models have a practical input limit.
# Chunk boundaries here are conservative (character-based, not token-based)
# to keep this dependency-free - swap for a tokenizer-aware chunker later
# if a specific model's limit demands more precision.
CHUNK_SIZE_CHARS = 4000
CHUNK_OVERLAP_CHARS = 200

# Below this character count for a document with pages, assume it's a
# scanned image with no real text layer rather than genuinely empty.
MIN_CHARS_PER_PAGE_FOR_NATIVE_TEXT = 20


@dataclass
class ProcessedDocument:
    full_text: str
    chunks: list[str]
    page_count: int
    needs_ocr: bool          # True if no meaningful native text layer was found
    file_format: str


def process_document(file_path: str) -> ProcessedDocument:
    extension = Path(file_path).suffix.lower().lstrip(".")

    if extension == "pdf":
        full_text, page_count = _extract_pdf_text(file_path)
    elif extension in ("docx", "doc"):
        full_text, page_count = _extract_docx_text(file_path)
    elif extension == "txt":
        full_text, page_count = _extract_txt_text(file_path)
    else:
        raise ValueError(f"Unsupported document format: {extension}")

    avg_chars_per_page = len(full_text) / page_count if page_count else len(full_text)
    needs_ocr = avg_chars_per_page < MIN_CHARS_PER_PAGE_FOR_NATIVE_TEXT

    chunks = _chunk_text(full_text) if full_text else []

    return ProcessedDocument(
        full_text=full_text,
        chunks=chunks,
        page_count=page_count,
        needs_ocr=needs_ocr,
        file_format=extension,
    )


def _extract_pdf_text(file_path: str) -> tuple[str, int]:
    text_parts = []
    with pdfplumber.open(file_path) as pdf:
        page_count = len(pdf.pages)
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            text_parts.append(page_text)
    return "\n".join(text_parts), page_count


def _extract_docx_text(file_path: str) -> tuple[str, int]:
    doc = DocxDocument(file_path)
    text_parts = [p.text for p in doc.paragraphs if p.text.strip()]
    full_text = "\n".join(text_parts)
    # DOCX has no fixed "page" concept without rendering - approximate
    # using paragraph count as a rough proxy, good enough for the
    # needs_ocr heuristic above (DOCX rarely needs OCR fallback anyway).
    approx_pages = max(1, len(text_parts) // 40)
    return full_text, approx_pages


def _extract_txt_text(file_path: str) -> tuple[str, int]:
    text = Path(file_path).read_text(encoding="utf-8", errors="replace")
    return text, 1


def _chunk_text(text: str) -> list[str]:
    if len(text) <= CHUNK_SIZE_CHARS:
        return [text]

    chunks = []
    start = 0
    while start < len(text):
        end = start + CHUNK_SIZE_CHARS
        chunks.append(text[start:end])
        start = end - CHUNK_OVERLAP_CHARS  # small overlap avoids cutting sentences/context at boundaries
    return chunks


if __name__ == "__main__":
    import sys

    path = sys.argv[1] if len(sys.argv) > 1 else "sample.pdf"
    result = process_document(path)

    print(f"Format: {result.file_format}")
    print(f"Pages: {result.page_count}")
    print(f"Total chars extracted: {len(result.full_text)}")
    print(f"Needs OCR fallback: {result.needs_ocr}")
    print(f"Chunks: {len(result.chunks)}")
    if result.chunks:
        print(f"First chunk preview: {result.chunks[0][:200]}...")